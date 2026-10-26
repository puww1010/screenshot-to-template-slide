#!/usr/bin/env python3
"""Regression tests for problems found in real use. Run:  python3 tests/test_regressions.py
Needs python-pptx (only to generate throw-away templates) and the same tools build_slide.py needs."""
import os, re, subprocess, sys, tempfile, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, 'scripts'); sys.path.insert(0, SCRIPTS)
import slidekit as sk  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

TMP = tempfile.mkdtemp(prefix='sts_test_')
FAILS = []


def check(name, cond, detail=''):
    print(('PASS ' if cond else 'FAIL ') + name + ('' if cond else f'  -> {detail}'))
    if not cond: FAILS.append(name)


def template(w, h, name):
    p = Presentation(); p.slide_width = Inches(w); p.slide_height = Inches(h)
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = 'x'
    path = os.path.join(TMP, name); p.save(path); return path


def page(body, name):
    path = os.path.join(TMP, name)
    open(path, 'w', encoding='utf-8').write(f'import sys; sys.path.insert(0, {SCRIPTS!r})\nfrom slidekit import *\n' + body)
    return path


def build(tpl, pg, *extra):
    out = os.path.join(TMP, os.path.basename(pg) + '.pptx')
    r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'build_slide.py'), '--template', tpl, '--page', pg,
                        '--out', out, '--pptx-skill', 'none', *extra], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr, out


def slide_xml(out):
    return zipfile.ZipFile(out).read('ppt/slides/slide1.xml').decode('utf-8')


T10 = template(10, 5.625, 't10.pptx')
T20 = template(20, 11.25, 't20.pptx')

# 1. para(lnpct): percent in, OOXML units (1/1000 %) out; OOXML-unit input still accepted; nonsense rejected.
#    Bug: lnpct=120 was written as spcPct val="120" (= 0.12 %), PowerPoint collapsed every line onto one.
check('para lnpct=120 -> 120000', 'spcPct val="120000"' in sk.para([sk.run('a', 12)], lnpct=120))
check('para lnpct=110000 kept', 'spcPct val="110000"' in sk.para([sk.run('a', 12)], lnpct=110000))
for bad in (5, 400, 900000):
    try:
        sk.para([sk.run('a', 12)], lnpct=bad); check(f'para lnpct={bad} rejected', False, 'no error raised')
    except ValueError:
        check(f'para lnpct={bad} rejected', True)

# 2. XML lint: spacing/size values PowerPoint takes literally (LibreOffice silently fixes them) must fail the build.
raw = page("def build():\n    return [text('T',.5,1,9,1,['<a:p><a:pPr><a:lnSpc><a:spcPct val=\"120\"/></a:lnSpc></a:pPr>'+run('one two',12)+'</a:p>'])]\n", 'bad_lnspc.py')
code, log, _ = build(T10, raw)
check('lint fails on spcPct val=120', code != 0 and '[FAIL] lint' in log, log[-300:])
raw = page("def build():\n    return [text('T',.5,1,9,1,[para([run('big',12).replace('sz=\"1200\"','sz=\"180000\"')])])]\n", 'bad_sz.py')
code, log, _ = build(T10, raw)
check('lint fails on sz=180000 (1800 pt)', code != 0 and '[FAIL] lint' in log, log[-300:])
ok = page("def build():\n    return [title_ph('Hello'), text('T',.5,1,9,1,[para([run('fine',12)],lnpct=115)])]\n", 'ok.py')
code, log, _ = build(T10, ok)
check('lint passes a normal page', code == 0 and '[lint] OK' in log, log[-300:])

# 3. --lang zh: CJK runs (including the title placeholder) are tagged zh-CN, Latin-only runs keep en-US.
zh = page("def build():\n    return [title_ph('标题 Title'), text('T',.5,1,9,1,[para([run('中文内容',12), run('English',12)])])]\n", 'zh.py')
code, log, out = build(T10, zh, '--lang', 'zh')
x = slide_xml(out) if code == 0 else ''
cjk_runs = re.findall(r'<a:r><a:rPr\b([^>]*)>(?:(?!</a:r>).)*?<a:t[^>]*>([^<]*[一-鿿][^<]*)</a:t>', x, re.S)
check('--lang zh tags every CJK run zh-CN', code == 0 and cjk_runs and all('lang="zh-CN"' in a for a, _ in cjk_runs), log[-300:])
check('--lang zh leaves Latin-only runs en-US', 'lang="en-US"' in x)
code, log, _ = build(T10, zh)
check('CJK text without --lang zh -> WARN', code == 0 and '[WARN]' in log and 'CJK' in log, log[-300:])

# 4. Canvas other than 10" x 5.625": slidekit coordinates/font sizes need scaling -> WARN with the factor.
code, log, _ = build(T20, ok)
check('20" canvas -> WARN with scale x2.0', code == 0 and re.search(r'\[WARN\] canvas 20(\.0+)? x 11\.25', log) and 'x2.0' in log, log[-400:])
# attribute order cy before cx (Google Slides exports write it this way) must be detected too
import shutil
T20b = os.path.join(TMP, 't20_cy_first.pptx'); d = os.path.join(TMP, 'unz'); shutil.rmtree(d, True)
zipfile.ZipFile(T20).extractall(d); pp = os.path.join(d, 'ppt/presentation.xml'); px = open(pp, encoding='utf-8').read()
px = re.sub(r'<p:sldSz cx="(\d+)" cy="(\d+)"', r'<p:sldSz cy="\2" cx="\1"', px); open(pp, 'w', encoding='utf-8').write(px)
subprocess.run(['zip', '-qXDr', T20b, '.'], cwd=d, check=True)
code, log, _ = build(T20b, ok)
check('20" canvas with cy before cx -> WARN', code == 0 and '[WARN] canvas 20' in log, log[-400:])
code, log, _ = build(T10, ok)
check('10" canvas -> no canvas WARN', '[WARN] canvas' not in log, log[-300:])

# 5. Fonts that are not standard on Windows/macOS -> NOTE to install/embed on the presenting machine.
code, log, _ = build(T10, ok, '--font', 'Figtree')
check('non-system font -> NOTE', code == 0 and '[NOTE]' in log and 'Figtree' in log, log[-300:])
code, log, _ = build(T10, ok, '--font', 'Arial')
check('Arial -> no font NOTE', '[NOTE] font' not in log, log[-300:])

print(f'\n{len(FAILS)} failed' if FAILS else '\nall passed')
sys.exit(1 if FAILS else 0)
